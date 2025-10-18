"""
Excel import/export endpoints backed by MongoDB
"""
from io import BytesIO
from typing import Any, Dict, List

from fastapi import APIRouter, UploadFile, File, HTTPException, status
from fastapi.responses import StreamingResponse
from openpyxl import load_workbook, Workbook

from app.core.config import settings
from app.core.mongodb import MongoDB


router = APIRouter(prefix="/files", tags=["Files"])


@router.post("/import", status_code=status.HTTP_200_OK)
async def import_excel(file: UploadFile = File(...)) -> Dict[str, Any]:
    """Import .xlsx: first row is header; insert rows into Mongo collection 'excel_data'."""
    filename = (file.filename or '').lower()
    if not filename.endswith('.xlsx'):
        raise HTTPException(status_code=400, detail="Only .xlsx files are supported")

    try:
        content = await file.read()
    finally:
        await file.close()

    try:
        wb = load_workbook(filename=BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="Failed to read Excel file")

    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {"inserted": 0}

    header = rows[0]
    if not header or all(h is None for h in header):
        raise HTTPException(status_code=400, detail="Missing header row")

    docs: List[Dict[str, Any]] = []
    for row in rows[1:]:
        if row is None:
            continue
        doc: Dict[str, Any] = {}
        for k, v in zip(header, row):
            if k is None:
                continue
            key = str(k).strip()
            doc[key] = v
        if any(v is not None and v != "" for v in doc.values()):
            docs.append(doc)

    if not docs:
        return {"inserted": 0}

    if MongoDB.client is None:
        raise HTTPException(status_code=500, detail="MongoDB is not initialized")
    db = MongoDB.client[settings.MONGODB_DB_NAME]
    result = await db["excel_data"].insert_many(docs)
    return {"inserted": len(result.inserted_ids)}


@router.get("/export")
async def export_excel() -> StreamingResponse:
    """Export all docs from Mongo 'excel_data' collection as .xlsx file."""
    if MongoDB.client is None:
        raise HTTPException(status_code=500, detail="MongoDB is not initialized")
    db = MongoDB.client[settings.MONGODB_DB_NAME]
    cursor = db["excel_data"].find({})

    items: List[Dict[str, Any]] = []
    async for doc in cursor:
        d = {k: v for k, v in doc.items() if k != "_id"}
        items.append(d)

    wb = Workbook()
    ws = wb.active
    ws.title = "data"

    if items:
        keys: List[str] = sorted({k for it in items for k in it.keys()})
        ws.append(keys)
        for it in items:
            ws.append([it.get(k) for k in keys])

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="export.xlsx"'},
    )

