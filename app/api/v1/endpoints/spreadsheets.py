"""
Spreadsheet endpoints
"""
from typing import List
from io import BytesIO
from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File
from fastapi.responses import StreamingResponse
from openpyxl import load_workbook, Workbook

from app.models import User
from app.api.v1.endpoints.auth import get_current_active_user
from app.services.spreadsheet_service import SpreadsheetService
from app.spreadsheet_schemas import (
    SpreadsheetCreate, SpreadsheetUpdate, SpreadsheetInfo,
    SpreadsheetDetail, SpreadsheetList, ShareSpreadsheet,
    CellUpdate, UpdatePermission
)
from app.mongodb_models import Spreadsheet, PermissionLevel


router = APIRouter(prefix="/spreadsheets", tags=["spreadsheets"])


@router.post("/", response_model=SpreadsheetInfo, status_code=status.HTTP_201_CREATED)
async def create_spreadsheet(
    spreadsheet_data: SpreadsheetCreate,
    current_user: User = Depends(get_current_active_user)
):
    """Create a new spreadsheet"""
    spreadsheet = await SpreadsheetService.create_spreadsheet(spreadsheet_data, current_user)
    return SpreadsheetInfo(
        id=str(spreadsheet.id),
        title=spreadsheet.title,
        owner_id=spreadsheet.owner_id,
        owner_email=spreadsheet.owner_email,
        rows=spreadsheet.rows,
        cols=spreadsheet.cols,
        created_at=spreadsheet.created_at,
        updated_at=spreadsheet.updated_at,
        version=spreadsheet.version,
        my_permission="owner"
    )


@router.get("/", response_model=SpreadsheetList)
async def get_my_spreadsheets(
    current_user: User = Depends(get_current_active_user)
):
    """Get all spreadsheets accessible by current user"""
    spreadsheets = await SpreadsheetService.get_user_spreadsheets(current_user)
    return SpreadsheetList(
        spreadsheets=spreadsheets,
        total=len(spreadsheets)
    )


@router.get("/{spreadsheet_id}", response_model=SpreadsheetDetail)
async def get_spreadsheet(
    spreadsheet_id: str,
    current_user: User = Depends(get_current_active_user)
):
    """Get spreadsheet details"""
    return await SpreadsheetService.get_spreadsheet(spreadsheet_id, current_user)


@router.patch("/{spreadsheet_id}", response_model=SpreadsheetInfo)
async def update_spreadsheet(
    spreadsheet_id: str,
    update_data: SpreadsheetUpdate,
    current_user: User = Depends(get_current_active_user)
):
    """Update spreadsheet metadata"""
    spreadsheet = await SpreadsheetService.update_spreadsheet(
        spreadsheet_id, update_data, current_user
    )
    permission = SpreadsheetService._get_user_permission(spreadsheet, current_user)
    return SpreadsheetInfo(
        id=str(spreadsheet.id),
        title=spreadsheet.title,
        owner_id=spreadsheet.owner_id,
        owner_email=spreadsheet.owner_email,
        rows=spreadsheet.rows,
        cols=spreadsheet.cols,
        created_at=spreadsheet.created_at,
        updated_at=spreadsheet.updated_at,
        version=spreadsheet.version,
        my_permission=permission
    )


@router.delete("/{spreadsheet_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_spreadsheet(
    spreadsheet_id: str,
    current_user: User = Depends(get_current_active_user)
):
    """Delete spreadsheet (owner only)"""
    await SpreadsheetService.delete_spreadsheet(spreadsheet_id, current_user)


@router.patch("/{spreadsheet_id}/cells")
async def update_cell(
    spreadsheet_id: str,
    cell_update: CellUpdate,
    current_user: User = Depends(get_current_active_user)
):
    """Update a single cell"""
    return await SpreadsheetService.update_cell(spreadsheet_id, cell_update, current_user)


@router.post("/{spreadsheet_id}/share", status_code=status.HTTP_200_OK)
async def share_spreadsheet(
    spreadsheet_id: str,
    share_data: ShareSpreadsheet,
    current_user: User = Depends(get_current_active_user)
):
    """Share spreadsheet with another user"""
    await SpreadsheetService.share_spreadsheet(spreadsheet_id, share_data, current_user)
    return {"message": f"Spreadsheet shared with {share_data.user_email}"}


@router.delete("/{spreadsheet_id}/permissions/{user_email}", status_code=status.HTTP_204_NO_CONTENT)
async def remove_permission(
    spreadsheet_id: str,
    user_email: str,
    current_user: User = Depends(get_current_active_user)
):
    """Remove user's access to spreadsheet"""
    await SpreadsheetService.remove_permission(spreadsheet_id, user_email, current_user)


# ===== Excel import/export =====

@router.post("/import", response_model=SpreadsheetInfo, status_code=status.HTTP_201_CREATED)
async def import_spreadsheet_from_excel(
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_active_user)
):
    """
    Import an .xlsx file and create a Spreadsheet owned by current user.
    The first non-empty sheet is used. Rows/cols are detected from data range.
    """
    filename = (file.filename or '').rsplit('/', 1)[-1]
    if not filename.lower().endswith('.xlsx'):
        raise HTTPException(status_code=400, detail="Only .xlsx files are supported")

    content = await file.read()
    await file.close()
    try:
        wb = load_workbook(filename=BytesIO(content), data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="Failed to read Excel file")

    # Pick the first sheet with any data
    ws = None
    for sheet in wb.worksheets:
        if sheet.max_row > 1 or sheet.max_column > 1:
            ws = sheet
            break
    if ws is None:
        ws = wb.active

    max_row = ws.max_row or 0
    max_col = ws.max_column or 0

    # Build cells map with zero-based indices keys: "r_c"
    cells: dict = {}
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            v = ws.cell(row=r, column=c).value
            if v is None or v == "":
                continue
            key = f"{r-1}_{c-1}"
            value_type = 'number' if isinstance(v, (int, float)) else 'boolean' if isinstance(v, bool) else 'date' if hasattr(v, 'isoformat') else 'text'
            # Formula detection is tricky with data_only=True; keep as text if values computed.
            cells[key] = {
                "value": v,
                "value_type": value_type,
                "formula": None,
                "style": {},
            }

    title = filename[:-5] if filename.lower().endswith('.xlsx') else filename or 'Imported Spreadsheet'
    spreadsheet = Spreadsheet(
        title=title or 'Imported Spreadsheet',
        owner_id=current_user.id,
        owner_email=current_user.email,
        rows=max(max_row, 100),
        cols=max(max_col, 26),
        cells=cells,
        permissions=[],
    )
    await spreadsheet.insert()

    return SpreadsheetInfo(
        id=str(spreadsheet.id),
        title=spreadsheet.title,
        owner_id=spreadsheet.owner_id,
        owner_email=spreadsheet.owner_email,
        rows=spreadsheet.rows,
        cols=spreadsheet.cols,
        created_at=spreadsheet.created_at,
        updated_at=spreadsheet.updated_at,
        version=spreadsheet.version,
        my_permission=PermissionLevel.OWNER,
    )


@router.get("/{spreadsheet_id}/export")
async def export_spreadsheet_to_excel(
    spreadsheet_id: str,
    current_user: User = Depends(get_current_active_user)
):
    """Export a spreadsheet (if user has access) as .xlsx."""
    detail = await SpreadsheetService.get_spreadsheet(spreadsheet_id, current_user)

    wb = Workbook()
    ws = wb.active
    ws.title = detail.title[:31] if detail.title else 'Sheet1'

    # Ensure at least 1x1 to avoid empty workbook
    max_r = max(detail.rows, 1)
    max_c = max(detail.cols, 1)

    # Fill values from cells map
    for key, cell in (detail.cells or {}).items():
        try:
            r_str, c_str = key.split('_', 1)
            r = int(r_str)
            c = int(c_str)
        except Exception:
            continue
        if r < 0 or c < 0:
            continue
        ws.cell(row=r+1, column=c+1, value=cell.get('value'))

    # Make sure sheet has proper dimensions (optional)
    ws.cell(row=max_r, column=max_c)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    filename = f"{detail.title or 'export'}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
